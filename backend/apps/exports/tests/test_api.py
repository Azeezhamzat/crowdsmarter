import io
import json
import zipfile
from datetime import timedelta

import pytest
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from apps.invitations.models import OrganisationInvitation
from apps.organisations.models import Membership


@pytest.mark.django_db
def test_owner_can_download_complete_export_without_secrets(
    api_client, user_factory, organisation_factory, decision_factory
):  # type: ignore[no-untyped-def]
    owner = user_factory(email="owner@example.com")
    organisation = organisation_factory(owner=owner)
    decision_factory(workspace=organisation.workspaces.get(is_default=True), owner=owner)
    OrganisationInvitation.objects.create(
        organisation=organisation,
        email="invitee@example.com",
        role=Membership.Role.CONTRIBUTOR,
        token_digest="a" * 64,
        invited_by=owner,
        expires_at=timezone.now() + timedelta(days=7),
    )
    api_client.force_authenticate(owner)

    response = api_client.get(
        reverse("exports:organisation-complete", kwargs={"organisation_id": organisation.id})
    )
    assert response.status_code == 200
    assert response["Content-Type"] == "application/zip"
    archive_bytes = bytes(response.content)
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        names = set(archive.namelist())
        assert "manifest.json" in names
        assert "json/decisions.json" in names
        assert "csv/evidence.csv" in names
        assert "json/invitations.json" in names
        assert "token_digest" not in archive.read("json/invitations.json").decode()
        assert "password" not in archive.read("json/users.json").decode()


@pytest.mark.django_db
def test_organisation_export_succeeds_with_the_browser_accept_header(
    api_client, user_factory, organisation_factory
):  # type: ignore[no-untyped-def]
    """The real frontend sends Accept: application/zip (see downloadApiFile in
    lib/api.ts). A prior regression only surfaced with this exact header: DRF's
    default content negotiation rejected it with 406 before the view ever ran,
    since APIClient.get() without an explicit Accept header trivially matches
    JSONRenderer and masks the bug."""
    owner = user_factory(email="owner2@example.com")
    organisation = organisation_factory(owner=owner)
    api_client.force_authenticate(owner)

    response = api_client.get(
        reverse("exports:organisation-complete", kwargs={"organisation_id": organisation.id}),
        HTTP_ACCEPT="application/zip",
    )

    assert response.status_code == 200
    assert response["Content-Type"] == "application/zip"


@pytest.mark.django_db
def test_organisation_export_includes_xlsx_workbook_and_stable_content_hash(
    api_client, user_factory, organisation_factory, decision_factory
):  # type: ignore[no-untyped-def]
    from openpyxl import load_workbook

    owner = user_factory(email="owner@example.com")
    organisation = organisation_factory(owner=owner)
    decision_factory(
        workspace=organisation.workspaces.get(is_default=True),
        owner=owner,
        title="Repeatable export decision",
    )
    api_client.force_authenticate(owner)
    url = reverse("exports:organisation-complete", kwargs={"organisation_id": organisation.id})

    first = api_client.get(url)
    second = api_client.get(url)
    assert first.status_code == 200 and second.status_code == 200

    with zipfile.ZipFile(io.BytesIO(bytes(first.content))) as archive:
        assert "xlsx/export.xlsx" in archive.namelist()
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["content_sha256"]) == 64
        first_hash = manifest["content_sha256"]

        workbook = load_workbook(io.BytesIO(archive.read("xlsx/export.xlsx")))
        assert "decisions" in workbook.sheetnames
        sheet = workbook["decisions"]
        headers = [cell.value for cell in sheet[1]]
        assert "title" in headers
        title_column = headers.index("title") + 1
        titles = [
            row[0].value
            for row in sheet.iter_rows(min_row=2, min_col=title_column, max_col=title_column)
        ]
        assert "Repeatable export decision" in titles

    with zipfile.ZipFile(io.BytesIO(bytes(second.content))) as archive:
        second_hash = json.loads(archive.read("manifest.json"))["content_sha256"]
    assert first_hash == second_hash

    decision_factory(
        workspace=organisation.workspaces.get(is_default=True),
        owner=owner,
        title="A newly added decision",
    )
    third = api_client.get(url)
    with zipfile.ZipFile(io.BytesIO(bytes(third.content))) as archive:
        third_hash = json.loads(archive.read("manifest.json"))["content_sha256"]
    assert third_hash != first_hash


@pytest.mark.django_db
def test_contributor_cannot_download_complete_organisation_export(
    api_client, user_factory, organisation_factory
):  # type: ignore[no-untyped-def]
    owner = user_factory()
    contributor = user_factory()
    organisation = organisation_factory(owner=owner)
    Membership.objects.create(
        organisation=organisation,
        user=contributor,
        role=Membership.Role.CONTRIBUTOR,
    )
    api_client.force_authenticate(contributor)
    response = api_client.get(
        reverse("exports:organisation-complete", kwargs={"organisation_id": organisation.id})
    )
    assert response.status_code == 403


@pytest.mark.django_db
def test_visible_member_can_download_decision_dossier(api_client, user_factory, decision_factory):  # type: ignore[no-untyped-def]
    decision = decision_factory(
        decision_question="Which option should we select?",
        purpose="Preserve the complete reasoning record.",
    )
    from apps.foresight.models import Source
    from apps.foresight.services import (
        create_research_claim,
        create_source,
        link_source_to_research_claim,
    )

    research_source = create_source(
        actor=decision.owner,
        organisation=decision.organisation,
        title="Decision research",
        source_type=Source.SourceType.RESEARCH,
        reference="Research-1",
    )
    claim = create_research_claim(
        actor=decision.owner,
        organisation=decision.organisation,
        statement="This decision requires a reversible implementation path.",
        linked_decision_id=decision.id,
    )
    link_source_to_research_claim(
        actor=decision.owner,
        claim=claim,
        source_id=research_source.id,
        relationship="supports",
    )
    viewer = user_factory()
    Membership.objects.create(
        organisation=decision.organisation,
        user=viewer,
        role=Membership.Role.VIEWER,
    )
    api_client.force_authenticate(viewer)
    response = api_client.get(
        reverse("exports:decision-dossier", kwargs={"decision_id": decision.id})
    )
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(bytes(response.content))) as archive:
        assert "summary.txt" in archive.namelist()
        assert "json/decision.json" in archive.namelist()
        assert decision.title in archive.read("summary.txt").decode()
        assert "xlsx/export.xlsx" in archive.namelist()
        assert "json/foresight_research_claims.json" in archive.namelist()
        assert "json/foresight_research_claim_sources.json" in archive.namelist()
        assert claim.statement in archive.read("json/foresight_research_claims.json").decode()
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["content_sha256"]) == 64

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(archive.read("xlsx/export.xlsx")))
        assert "options" in workbook.sheetnames


@pytest.mark.django_db
def test_decision_export_succeeds_with_the_browser_accept_header(
    api_client, user_factory, decision_factory
):  # type: ignore[no-untyped-def]
    decision = decision_factory()
    api_client.force_authenticate(decision.owner)

    response = api_client.get(
        reverse("exports:decision-dossier", kwargs={"decision_id": decision.id}),
        HTTP_ACCEPT="application/zip",
    )

    assert response.status_code == 200
    assert response["Content-Type"] == "application/zip"


@pytest.mark.django_db
def test_outsider_cannot_discover_decision_through_export(
    api_client, user_factory, decision_factory
):  # type: ignore[no-untyped-def]
    decision = decision_factory()
    outsider = user_factory()
    api_client.force_authenticate(outsider)
    response = api_client.get(
        reverse("exports:decision-dossier", kwargs={"decision_id": decision.id})
    )
    assert response.status_code == 404


@pytest.mark.django_db
def test_organisation_export_contains_foresight_metadata_and_private_file(
    api_client, organisation_factory
):  # type: ignore[no-untyped-def]
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.foresight.services import (
        attach_source_file,
        create_research_claim,
        create_source,
        link_source_to_research_claim,
    )

    organisation = organisation_factory()
    owner = organisation.created_by
    source = create_source(
        actor=owner,
        organisation=organisation,
        title="Emerging regulation report",
        source_type="government",
        reference="REG-2026-01",
        credibility="high",
        credibility_rationale="Published by the responsible regulator.",
    )
    attachment = attach_source_file(
        actor=owner,
        source=source,
        upload=SimpleUploadedFile(
            "regulation.pdf",
            b"%PDF-1.4 simulated regulation",
            content_type="application/pdf",
        ),
    )
    claim = create_research_claim(
        actor=owner,
        organisation=organisation,
        statement="The regulation changes the decision environment.",
        state="supported",
        authority_score=3,
        directness_score=3,
        recency_score=2,
        triangulation_score=1,
    )
    link_source_to_research_claim(
        actor=owner,
        claim=claim,
        source_id=source.id,
        relationship="supports",
        note="The responsible regulator issued the source.",
    )
    api_client.force_authenticate(owner)

    response = api_client.get(
        reverse("exports:organisation-complete", kwargs={"organisation_id": organisation.id})
    )

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(bytes(response.content))) as archive:
        names = set(archive.namelist())
        assert "json/foresight_sources.json" in names
        assert "json/foresight_source_attachments.json" in names
        assert "json/foresight_research_claims.json" in names
        assert "csv/foresight_research_claims.csv" in names
        assert "json/foresight_research_claim_sources.json" in names
        assert any(
            name.startswith(f"attachments/foresight/{source.id}/{attachment.id}.") for name in names
        )


@pytest.mark.django_db
@override_settings(SOURCE_ATTACHMENT_ENFORCE_CLEAN_DOWNLOADS=True)
def test_organisation_export_omits_attachment_without_clean_scan(api_client, organisation_factory):  # type: ignore[no-untyped-def]
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.foresight.models import SourceAttachment
    from apps.foresight.services import attach_source_file, create_source

    organisation = organisation_factory()
    owner = organisation.created_by
    source = create_source(
        actor=owner,
        organisation=organisation,
        title="Quarantined report",
        source_type="internal",
        reference="QUARANTINE-1",
    )
    attachment = attach_source_file(
        actor=owner,
        source=source,
        upload=SimpleUploadedFile("report.txt", b"Initially clean", content_type="text/plain"),
    )
    SourceAttachment.objects.filter(id=attachment.id).update(
        malware_scan_status=SourceAttachment.ScanStatus.INFECTED
    )
    api_client.force_authenticate(owner)

    response = api_client.get(
        reverse("exports:organisation-complete", kwargs={"organisation_id": organisation.id})
    )

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(bytes(response.content))) as archive:
        attachment_entries = [
            name
            for name in archive.namelist()
            if name.startswith(f"attachments/foresight/{source.id}/{attachment.id}.")
        ]
        assert attachment_entries == [
            f"attachments/foresight/{source.id}/{attachment.id}.unavailable.txt"
        ]
        assert b"not been verified as malware-free" in archive.read(attachment_entries[0])


@pytest.mark.django_db
def test_organisation_export_preserves_systems_mapping_and_feedback_loop_order(
    api_client, organisation_factory
):  # type: ignore[no-untyped-def]
    from apps.foresight.mapping_services import (
        create_canvas,
        create_driver,
        create_feedback_loop,
    )

    organisation = organisation_factory()
    owner = organisation.created_by
    canvas = create_canvas(
        actor=owner,
        organisation=organisation,
        title="Adaptive capacity system",
        focal_question="How does adaptive capacity reinforce future resilience?",
        scope="The organisation and its implementation ecosystem.",
        horizon_year=2035,
    )
    capability = create_driver(
        actor=owner,
        canvas=canvas,
        title="Internal adaptive capability",
        description="Capability affects the speed and quality of organisational adaptation.",
        driver_type="driver",
        steep_category="social",
        impact=4,
        uncertainty=3,
    )
    confidence = create_driver(
        actor=owner,
        canvas=canvas,
        title="Stakeholder confidence",
        description="Confidence influences willingness to invest in further adaptation.",
        driver_type="driver",
        steep_category="economic",
        impact=4,
        uncertainty=4,
    )
    loop = create_feedback_loop(
        actor=owner,
        canvas=canvas,
        name="Adaptive confidence loop",
        description="Capability builds confidence, which enables further capability investment.",
        loop_type="reinforcing",
        driver_ids=[confidence.id, capability.id],
        rationale="The loop captures the team's explicit interpretation of the system.",
    )
    api_client.force_authenticate(owner)

    response = api_client.get(
        reverse(
            "exports:organisation-complete",
            kwargs={"organisation_id": organisation.id},
        )
    )

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(bytes(response.content))) as archive:
        names = set(archive.namelist())
        assert "json/foresight_canvases.json" in names
        assert "json/foresight_drivers.json" in names
        assert "json/foresight_feedback_loops.json" in names
        assert "json/foresight_feedback_loop_drivers.json" in names
        links = json.loads(archive.read("json/foresight_feedback_loop_drivers.json").decode())
        ordered = sorted(
            (item for item in links if item["feedback_loop"] == str(loop.id)),
            key=lambda item: item["position"],
        )
        assert [item["driver"] for item in ordered] == [
            str(confidence.id),
            str(capability.id),
        ]


@pytest.mark.django_db
def test_organisation_export_includes_open_session_ideation_data_without_token_digest(
    organisation_factory,
):  # type: ignore[no-untyped-def]
    from apps.exports.services import build_organisation_export
    from apps.ideation import services as ideation_services

    organisation = organisation_factory()
    owner = organisation.created_by
    session = ideation_services.create_session(
        actor=owner,
        organisation=organisation,
        title="Ideathon",
        prompt="What should we try?",
    )
    session = ideation_services.open_session(actor=owner, session=session)
    participant, _token = ideation_services.identify_participant(
        session=session,
        name="External participant",
        email="participant@example.com",
    )
    idea = ideation_services.submit_idea(
        session=session,
        participant=participant,
        title="A promising idea",
        description="Details.",
    )
    ideation_services.cast_vote(idea=idea, participant=participant)

    archive = build_organisation_export(organisation=organisation)
    with zipfile.ZipFile(io.BytesIO(archive.content)) as zf:
        names = set(zf.namelist())
        assert "json/open_sessions.json" in names
        assert "json/open_session_participants.json" in names
        assert "json/ideas.json" in names
        assert "json/idea_votes.json" in names

        sessions = json.loads(zf.read("json/open_sessions.json").decode())
        assert [s["title"] for s in sessions] == ["Ideathon"]

        ideas = json.loads(zf.read("json/ideas.json").decode())
        assert [i["title"] for i in ideas] == ["A promising idea"]

        participants_payload = zf.read("json/open_session_participants.json").decode()
        assert "token_digest" not in participants_payload
        assert "participant@example.com" in participants_payload


@pytest.mark.django_db
def test_decision_export_scopes_open_sessions_to_that_decision(
    decision_factory,
):  # type: ignore[no-untyped-def]
    from apps.exports.services import build_decision_export
    from apps.ideation import services as ideation_services

    decision = decision_factory()
    owner = decision.organisation.created_by
    linked_session = ideation_services.create_session(
        actor=owner,
        organisation=decision.organisation,
        decision=decision,
        title="Community input",
        prompt="What matters here?",
    )
    ideation_services.create_session(
        actor=owner,
        organisation=decision.organisation,
        title="Unrelated ideathon",
        prompt="Something else entirely",
    )

    archive = build_decision_export(decision=decision)
    with zipfile.ZipFile(io.BytesIO(archive.content)) as zf:
        sessions = json.loads(zf.read("json/open_sessions.json").decode())
        assert [s["title"] for s in sessions] == [linked_session.title]


@pytest.mark.django_db
def test_grant_round_export_includes_budget_summary_and_award_letters(
    decision_factory,
):  # type: ignore[no-untyped-def]
    from apps.decision_options.models import DecisionOption
    from apps.decision_options.services import create_option, set_outcome
    from apps.exports.services import build_decision_export
    from apps.ideation import services as ideation_services

    decision = decision_factory(source_template_key="grant_round", title="2026 community round")
    owner = decision.owner
    session = ideation_services.create_session(
        actor=owner,
        organisation=decision.organisation,
        decision=decision,
        title="Round intake",
        prompt="Applications?",
    )
    session = ideation_services.open_session(actor=owner, session=session)
    participant, _token = ideation_services.identify_participant(
        session=session,
        name="Ada Lovelace",
        email="ada@example.com",
    )
    idea = ideation_services.submit_idea(
        session=session,
        participant=participant,
        title="Community garden expansion",
        description="Expand shared plots.",
        requested_amount="1000.00",
    )
    funded_option = ideation_services.promote_idea_to_decision(
        actor=owner, idea=idea, decision=decision
    )
    set_outcome(
        actor=owner,
        option=funded_option,
        outcome_status=DecisionOption.OutcomeStatus.FUNDED,
        awarded_amount="800.00",
        outcome_note="Congratulations.",
    )
    declined_option = create_option(
        actor=owner,
        decision=decision,
        title="Directly created application",
        description="Desc.",
        estimated_cost="500.00",
    )
    set_outcome(
        actor=owner, option=declined_option, outcome_status=DecisionOption.OutcomeStatus.DECLINED
    )

    archive = build_decision_export(decision=decision)
    with zipfile.ZipFile(io.BytesIO(archive.content)) as zf:
        summary = zf.read("summary.txt").decode()
        assert "Grant round budget" in summary
        assert "Total awarded: 800.00" in summary

        names = set(zf.namelist())
        letter_names = [n for n in names if n.startswith("letters/")]
        assert len(letter_names) == 2

        funded_letter = zf.read(
            f"letters/{funded_option.title.lower().replace(' ', '-')}.txt"
        ).decode()
        assert "Ada Lovelace" in funded_letter
        assert "800.00" in funded_letter
        assert "ada@example.com" in funded_letter

        declined_letter = zf.read(
            f"letters/{declined_option.title.lower().replace(' ', '-')}.txt"
        ).decode()
        assert "was not funded" in declined_letter


@pytest.mark.django_db
def test_decision_dossier_includes_printable_report_without_confidential_identity(
    decision_factory,
):  # type: ignore[no-untyped-def]
    from apps.contributions.services import (
        create_facilitation_record,
        create_session,
        update_agenda_item_status,
        update_session_status,
    )
    from apps.exports.services import build_decision_export

    decision = decision_factory(status="open_for_contribution", title="Access decision")
    session = create_session(
        actor=decision.owner,
        decision=decision,
        title="Access workshop",
        objective="Understand access barriers.",
        participation_channels=["phone"],
        agenda_items=[
            {
                "title": "Surface access barriers",
                "method": "Confidential telephone input",
                "planned_minutes": 20,
            }
        ],
    )
    session = update_session_status(actor=decision.owner, session=session, status="open")
    agenda_item = update_agenda_item_status(
        actor=decision.owner, item=session.agenda_items.get(), action="start"
    )
    create_facilitation_record(
        actor=decision.owner,
        session=session,
        kind="participant_statement",
        body="Evening access creates a safety concern.",
        channel="phone",
        origin="participant_input",
        attribution="confidential",
        speaker_label="Protected participant identity",
        agenda_item_id=agenda_item.id,
    )

    export = build_decision_export(decision=decision)
    with zipfile.ZipFile(io.BytesIO(export.content)) as archive:
        records = json.loads(archive.read("json/facilitation_records.json"))
        agenda_items = json.loads(archive.read("json/facilitation_agenda_items.json"))
        assert agenda_items[0]["title"] == "Surface access barriers"
        assert records[0]["speaker_label"] == ""
        assert records[0]["source_participant"] is None
        report_name = next(
            name for name in archive.namelist() if name.startswith("reports/facilitation-")
        )
        report = archive.read(report_name).decode("utf-8")
        assert "Evening access creates a safety concern." in report
        assert "Confidential participant input" in report
        assert "Surface access barriers" in report
        assert "Protected participant identity" not in report
